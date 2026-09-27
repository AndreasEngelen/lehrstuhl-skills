"""Shared helpers for Chair-style Word tables (python-docx).

Chair defaults: Times New Roman, three horizontal rules (top, below header,
bottom), no vertical lines, caption above the table, note below the table.
"""
from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, Cm

MINUS = "−"


def fmt_num(x, decimals=2, leading_zero=True, thousands=True):
    """Format a number with a true minus sign and fixed decimals."""
    if x is None or x == "":
        return ""
    if isinstance(x, str):
        return x
    s = f"{abs(x):,.{decimals}f}" if thousands else f"{abs(x):.{decimals}f}"
    if not leading_zero and s.startswith("0."):
        s = s[1:]
    if x < 0 and float(s.replace(",", "") or 0) != 0:
        s = MINUS + s
    return s


def new_document(landscape=False, font="Times New Roman", size=10):
    doc = Document()
    st = doc.styles["Normal"]
    st.font.name = font
    st.element.rPr.rFonts.set(qn("w:eastAsia"), font)
    st.font.size = Pt(size)
    st.paragraph_format.space_after = Pt(0)
    st.paragraph_format.space_before = Pt(0)
    st.paragraph_format.line_spacing = 1.0
    sec = doc.sections[0]
    if landscape:
        sec.orientation = WD_ORIENT.LANDSCAPE
        sec.page_width, sec.page_height = sec.page_height, sec.page_width
    for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(sec, side, Cm(2))
    return doc


def _border(cell, edge, sz=8):
    tcPr = cell._tc.get_or_add_tcPr()
    borders = tcPr.find(qn("w:tcBorders"))
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tcPr.append(borders)
    el = OxmlElement(f"w:{edge}")
    el.set(qn("w:val"), "single")
    el.set(qn("w:sz"), str(sz))
    el.set(qn("w:space"), "0")
    el.set(qn("w:color"), "000000")
    borders.append(el)


def rule_below(row, sz=8):
    for c in row.cells:
        _border(c, "bottom", sz)


def rule_above(row, sz=8):
    for c in row.cells:
        _border(c, "top", sz)


def set_cell(cell, text, bold=False, italic=False, align="center", size=None, superscript=None):
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.space_before = Pt(0)
    p.alignment = {"center": WD_ALIGN_PARAGRAPH.CENTER, "left": WD_ALIGN_PARAGRAPH.LEFT,
                   "right": WD_ALIGN_PARAGRAPH.RIGHT}[align]
    r = p.add_run(str(text))
    r.bold, r.italic = bold, italic
    if size:
        r.font.size = Pt(size)
    if superscript:
        s = p.add_run(superscript)
        s.font.superscript = True
        if size:
            s.font.size = Pt(size)
    return p


def add_caption(doc, number, title, label="Table"):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run(f"{label} {number}. ")
    r.bold = True
    p.add_run(title)
    return p


def add_note(doc, text, size=9):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(3)
    r = p.add_run("Note: ")
    r.italic = True
    r.font.size = Pt(size)
    r2 = p.add_run(text)
    r2.font.size = Pt(size)
    return p


def new_table(doc, rows, cols, first_col_cm=None, other_col_cm=None):
    """Create a table; optionally fix the width of the label column and the data columns."""
    t = doc.add_table(rows=rows, cols=cols)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    tblPr = t._tbl.tblPr
    mar = OxmlElement("w:tblCellMar")
    for side, w in (("left", 57), ("right", 57), ("top", 0), ("bottom", 0)):
        el = OxmlElement(f"w:{side}")
        el.set(qn("w:w"), str(w))
        el.set(qn("w:type"), "dxa")
        mar.append(el)
    tblPr.append(mar)
    if first_col_cm:
        t.autofit = False
        for j, col in enumerate(t.columns):
            col.width = Cm(first_col_cm if j == 0 else other_col_cm)
        for row in t.rows:
            for j, cell in enumerate(row.cells):
                cell.width = Cm(first_col_cm if j == 0 else other_col_cm)
    for row in t.rows:
        for cell in row.cells:
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.space_before = Pt(0)
    return t


def usable_width_cm(doc):
    sec = doc.sections[0]
    return (sec.page_width - sec.left_margin - sec.right_margin) / 360000
